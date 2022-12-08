import docx
import pandas as pd
import os


def read_output(file):
    with open(file, 'r') as f:
        filetext_list = f.readlines()
    return filetext_list

def read_numbers(file):
    list_of_numbers = []
    with open(file, 'r') as f:
        for line in f:
            try:
                list_of_numbers.append(int(line))
            except ValueError:
                pass
    return list_of_numbers

def write_doc(data, fname):
    doc = docx.Document() 
    t = doc.add_table(data.shape[0]+1, data.shape[1]) # add a table to the end and create a reference variable; extra row is so we can add the header row 
    
    for j in range(data.shape[-1]): # add the header rows
        t.cell(0,j).text = data.columns[j]

    for i in range(data.shape[0]):# add the rest of the data frame
        for j in range(data.shape[-1]):
            t.cell(i+1,j).text = "%4.4f" % (float((data.values[i,j])),)
    doc.save(fname) # save the doc


def CP_data(N, text_list):
    data_dict = {'N': N}
    block_title = "CP #%5d\n" % (N,)
    i = text_list.index(block_title)
    data_dict['rho'] = text_list[i+24].strip().split()[2]
    data_dict['Lap'] = text_list[i+29].strip().split()[2]
    data_dict['V'] = text_list[i+46].strip().split()[2]
    data_dict['G'] = text_list[i+45].strip().split()[2]
    data_dict['H'] = text_list[i+47].strip().split()[2]
    data_dict['M'] = text_list[i+32].strip().split()[2]
    data_dict['elip'] = text_list[i+33].strip().split()[2]
    return data_dict

if __name__ == '__main__':
    output_text_list = read_output("nohup.out")
    list_of_numbers = read_numbers("points.txt")
    data = pd.DataFrame(columns = ['N', 'rho', 'Lap', 'V', 'G', 'H', 'M', 'elip'])
    for i in list_of_numbers:
        data = data.append(CP_data(i, output_text_list), ignore_index = True)
        #print(text)
    result_table = data.to_string()
    print(result_table)
    write_doc(data, 'results.docx')

